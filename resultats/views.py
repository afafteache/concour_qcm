from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.contrib import messages
from accounts.decorators import role_required
from accounts.models import Profil
from .models import Resultat
import openpyxl
from openpyxl import load_workbook


@login_required(login_url='login')
@role_required('admin')
def resultats_liste(request):
    resultats = Resultat.objects.select_related('candidat', 'examen').all().order_by('-date_passage')

    total = resultats.count()
    valides = 0
    non_valides = 0
    stats = []

    for r in resultats:
        try:
            total_q = r.examen.questions.count()
            pourcentage = round((r.score / total_q * 100), 1) if total_q > 0 else 0
            reussi = pourcentage >= 50
        except Exception:
            pourcentage = 0
            reussi = False
            total_q = 0

        if reussi:
            valides += 1
        else:
            non_valides += 1

        profil = Profil.objects.filter(user=r.candidat).first()
        stats.append({
            'resultat': r,
            'candidat': r.candidat.get_full_name(),
            'cin': profil.cin if profil else '—',
            'examen': r.examen.titre,
            'score': r.score,
            'total': total_q,
            'pourcentage': pourcentage,
            'reussi': reussi,
            'date': r.date_passage,
        })

    return render(request, 'admin/resultats/liste.html', {
        'stats': stats,
        'total': total,
        'valides': valides,
        'non_valides': non_valides,
        'pct_valides': round((valides / total * 100), 1) if total > 0 else 0,
    })


@login_required(login_url='login')
@role_required('admin')
def import_candidats_excel(request):
    if request.method == 'POST' and request.FILES.get('fichier_excel'):
        fichier = request.FILES['fichier_excel']
        try:
            wb = load_workbook(fichier)
            ws = wb.active
            crees = 0
            erreurs = []
            doublons = 0

            for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not any(row):
                    continue
                try:
                    prenom = str(row[0]).strip() if row[0] else ''
                    nom    = str(row[1]).strip() if row[1] else ''
                    cin    = str(row[2]).strip() if row[2] else ''

                    if not prenom or not nom or not cin:
                        erreurs.append(f"Ligne {i} : données manquantes.")
                        continue

                    if Profil.objects.filter(cin=cin).exists():
                        doublons += 1
                        continue

                    user = User.objects.create_user(
                        username=cin,
                        password=User.objects.make_random_password(),
                        first_name=prenom,
                        last_name=nom,
                    )
                    Profil.objects.create(user=user, cin=cin, role='candidat')
                    crees += 1

                except Exception as e:
                    erreurs.append(f"Ligne {i} : erreur — {e}")

            msg = f"{crees} candidat(s) importé(s) avec succès."
            if doublons > 0:
                msg += f" {doublons} doublon(s) ignoré(s)."
            if erreurs:
                msg += f" {len(erreurs)} erreur(s)."

            if crees > 0:
                messages.success(request, msg)
            else:
                messages.warning(request, msg)

            for e in erreurs:
                messages.error(request, e)

        except Exception as e:
            messages.error(request, f"Erreur lors de la lecture du fichier : {e}")

        return redirect('import_candidats')

    candidats = Profil.objects.filter(role='candidat').select_related('user').order_by('user__last_name')
    return render(request, 'admin/candidats/import.html', {'candidats': candidats})


@login_required(login_url='login')
@role_required('admin')
def supprimer_candidat(request, profil_id):
    """Supprimer un candidat."""
    profil = get_object_or_404(Profil, id=profil_id, role='candidat')
    if request.method == 'POST':
        nom = profil.user.get_full_name()
        profil.user.delete()  # supprime aussi le profil (CASCADE)
        messages.success(request, f'Candidat {nom} supprimé avec succès.')
    return redirect('import_candidats')


@login_required(login_url='login')
@role_required('admin')
def telecharger_modele_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Candidats"
    ws['A1'] = 'Prénom'
    ws['B1'] = 'Nom'
    ws['C1'] = 'CIN'

    from openpyxl.styles import Font, PatternFill, Alignment
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1a2a5e", end_color="1a2a5e", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    ws['A2'] = 'Khadija'
    ws['B2'] = 'Daji'
    ws['C2'] = 'OD67855'

    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 15

    from django.http import HttpResponse
    import io
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="modele_candidats.xlsx"'
    return response